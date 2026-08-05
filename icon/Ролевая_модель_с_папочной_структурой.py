# -*- coding: utf-8 -*-
"""
SGNL / СОД Permissions + Folder Structure GUI v11
--------------------------------------------------
Утилита для создания папочной структуры и выдачи прав на папки
docs.sgnl.pro по одной Excel-матрице.

Файлы, под которые написана версия:
    1) Excel с листом "2. Ролевая матрица" или активным листом: колонки "Уровень N" + колонки ролей.
    2) Дизайн/UX взят по мотиву существующего VitroCAD_Permissions.py.

Зависимости:
    pip install requests openpyxl PySide6
Если используется PyQt5:
    pip install requests openpyxl PyQt5 PyQtWebEngine

Важно по авторизации:
    Основной способ — интерактивный OAuth PKCE во встроенном браузере.
    Это необходимо, потому что актуальный вход SGNL использует страницу
    auth.sgnl.pro и CAPTCHA. Hub вручную открывать не нужно, но сам
    hub.sgnl.pro остаётся OAuth-сервером и выдаёт access_token docs_web.
    Токен также можно вставить вручную для диагностики.

Логика прав по легенде Excel:
    - / пусто  -> не трогать эту роль на этой папке
    П          -> read
    С          -> read + download
    З          -> read + download + create
    Р          -> read + download + create + update + delete

Основные API-маршруты:
    GET  {docs}/api/folders/project/{projectId}/tree
    PUT  {docs}/api/folders
    GET  {docs}/api/permissions/tree?projectId={projectId}
    GET  {hub}/api/v1/hub/companies/{companyId}/custom/roles
    GET  {hub}/api/v1/hub/companies/{companyId}/users
    OAuth использует client_id docs_web и scopes из реального DOCS web-клиента; hub-запросы выполняются с browser-like заголовками.
    POST {docs}/api/permissions/folders/{folderId}/upsert
"""

from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import secrets
import sys
import traceback
import unicodedata
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import parse_qs, urlencode, urlparse

import requests
import urllib3
from openpyxl import load_workbook

WEBENGINE_AVAILABLE = False
QtWebEngineWidgets = None

try:
    from PySide6 import QtCore, QtGui, QtWidgets

    Signal = QtCore.Signal
    Slot = QtCore.Slot
    try:
        from PySide6 import QtWebEngineWidgets as _QtWebEngineWidgets

        QtWebEngineWidgets = _QtWebEngineWidgets
        WEBENGINE_AVAILABLE = True
    except ImportError:
        pass
except ImportError:  # pragma: no cover
    from PyQt5 import QtCore, QtGui, QtWidgets

    Signal = QtCore.pyqtSignal
    Slot = QtCore.pyqtSlot
    try:
        from PyQt5 import QtWebEngineWidgets as _QtWebEngineWidgets

        QtWebEngineWidgets = _QtWebEngineWidgets
        WEBENGINE_AVAILABLE = True
    except ImportError:
        pass


# -----------------------------------------------------------------------------
# Runtime defaults. Company/project are intentionally empty until login.
# -----------------------------------------------------------------------------
DEFAULT_DOCS_URL = "https://docs.sgnl.pro"
DEFAULT_HUB_URL = "https://hub.sgnl.pro"
DEFAULT_PROJECT_ID = ""
DEFAULT_COMPANY_ID = ""
DEFAULT_COMPANY_NAME = ""
DEFAULT_PROJECT_NAME = ""

OAUTH_CLIENT_ID = "docs_web"
OAUTH_SCOPE = (
    "offline_access cl.front cl.sys hub.r hub.e "
    "doc.c doc.r doc.u doc.d doc.e doc.x "
    "conv.x conv.r cnr.r cnr.e lg.c"
)

# Browser User-Agent for OAuth and API requests.
# It must exist before SgnlClient methods build browser-like headers.
BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

VERIFY_SSL_CERTIFICATE = True
if not VERIFY_SSL_CERTIFICATE:
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

RIGHT_ORDER = ["П", "С", "З", "Р"]
RIGHT_LABELS = {
    "П": "Просмотр",
    "С": "Скачивание",
    "З": "Загрузка",
    "Р": "Редактирование",
}
RIGHT_FLAGS: Dict[str, Dict[str, bool]] = {
    "П": {"read": True, "download": False, "create": False, "update": False, "delete": False},
    "С": {"read": True, "download": True, "create": False, "update": False, "delete": False},
    "З": {"read": True, "download": True, "create": True, "update": False, "delete": False},
    "Р": {"read": True, "download": True, "create": True, "update": True, "delete": True},
}
ZERO_FLAGS = {"read": False, "download": False, "create": False, "update": False, "delete": False}

RIGHT_ALIASES = {
    "п": "П",
    "p": "П",
    "read": "П",
    "view": "П",
    "просмотр": "П",
    "чтение": "П",
    "с": "С",
    "c": "С",
    "download": "С",
    "скачивание": "С",
    "скачать": "С",
    "з": "З",
    "z": "З",
    "upload": "З",
    "create": "З",
    "создание": "З",
    "загрузка": "З",
    "загрузить": "З",
    "р": "Р",
    "r": "Р",
    "edit": "Р",
    "update": "Р",
    "modify": "Р",
    "редактирование": "Р",
    "изменение": "Р",
    "изменить": "Р",
    "полный доступ": "Р",
}


# -----------------------------------------------------------------------------
# Text helpers
# -----------------------------------------------------------------------------
def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\r", " ").replace("\n", " ").strip()
    text = re.sub(r"\s+", " ", text)
    if text.casefold() in {"none", "nan", "null"}:
        return ""
    return text


def normalize_text(value: object) -> str:
    text = clean_text(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("ё", "е").replace("Ё", "Е")
    text = text.casefold()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_path_parts(parts: Sequence[object]) -> str:
    return "/".join(normalize_text(part) for part in parts if clean_text(part))


def display_path(parts: Sequence[object]) -> str:
    return " / ".join(clean_text(part) for part in parts if clean_text(part))


def best_close_matches(value: str, variants: Iterable[str], limit: int = 5, cutoff: float = 0.68) -> List[str]:
    import difflib

    norm_variants = list(variants)
    return difflib.get_close_matches(value, norm_variants, n=limit, cutoff=cutoff)


def bearer_header(token: str) -> str:
    token = clean_text(token)
    if not token:
        return ""
    if token.lower().startswith("bearer "):
        return token
    return "Bearer " + token


def decode_jwt_payload(token: str) -> dict:
    """Диагностически декодирует payload JWT без проверки подписи.

    Нужен только для лога: client_id/scopes сразу показывают,
    каким web-клиентом получен token.
    """
    token = clean_text(token)
    if token.lower().startswith("bearer "):
        token = token.split(" ", 1)[1].strip()
    parts = token.split(".")
    if len(parts) < 2:
        return {}
    payload = parts[1]
    payload += "=" * ((4 - len(payload) % 4) % 4)
    try:
        return json.loads(base64.urlsafe_b64decode(payload.encode("ascii")).decode("utf-8"))
    except Exception:
        return {}


def base64url_no_padding(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode("ascii").rstrip("=")


def make_pkce_pair() -> Tuple[str, str]:
    # RFC 7636: verifier 43..128 chars. token_urlsafe(64) is normally 86 chars.
    verifier = secrets.token_urlsafe(64)[:128]
    challenge = base64url_no_padding(hashlib.sha256(verifier.encode("ascii")).digest())
    return verifier, challenge



# -----------------------------------------------------------------------------
# Data models
# -----------------------------------------------------------------------------
@dataclass
class PermissionFlags:
    code: str
    flags: Dict[str, bool]

    @property
    def label(self) -> str:
        return RIGHT_LABELS.get(self.code, self.code)

    @property
    def text(self) -> str:
        if not self.code:
            return ""
        return f"{self.code} ({self.label})"


@dataclass
class PermissionEntry:
    row_number: int
    path_parts: List[str]
    permissions: Dict[str, PermissionFlags]

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)


@dataclass
class FolderItem:
    id: str
    name: str
    path_parts: List[str]
    permissions: List[dict] = field(default_factory=list)

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)

    @property
    def norm_path(self) -> str:
        return normalize_path_parts(self.path_parts)


@dataclass
class Principal:
    name: str
    id: str
    source_type: str = "Role"  # Role / User


@dataclass
class PrincipalResolution:
    requested_name: str
    principal: Optional[Principal] = None
    message: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.principal and self.principal.id)


@dataclass
class PlanRow:
    entry: PermissionEntry
    folder: Optional[FolderItem]
    status: str
    message: str
    principals: Dict[str, PrincipalResolution] = field(default_factory=dict)

    @property
    def can_apply(self) -> bool:
        return self.status == "OK" and self.folder is not None and all(res.ok for res in self.principals.values())

    @property
    def permission_count(self) -> int:
        return len(self.entry.permissions)


@dataclass
class FolderPathEntry:
    row_number: int
    path_parts: List[str]

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)


@dataclass
class FolderCreatePlanRow:
    row_number: int
    path_parts: List[str]
    status: str
    parent_path: str
    folder_id: str = ""
    message: str = ""

    @property
    def path_text(self) -> str:
        return display_path(self.path_parts)

    @property
    def name(self) -> str:
        return self.path_parts[-1] if self.path_parts else ""

    @property
    def can_create(self) -> bool:
        return self.status == "Создать"


# -----------------------------------------------------------------------------
# Excel parser
# -----------------------------------------------------------------------------
def right_from_cell(value: object) -> Optional[PermissionFlags]:
    text = clean_text(value)
    if not text:
        return None
    if text in {"-", "—", "–", "нет", "Нет", "0"}:
        return None

    # Keep one-letter Cyrillic codes intact, but also accept words / English aliases.
    tokens = [clean_text(part) for part in re.split(r"[,;\n\r/]+", text) if clean_text(part)]
    if not tokens:
        tokens = [text]

    codes: List[str] = []
    for token in tokens:
        norm = normalize_text(token)
        code = RIGHT_ALIASES.get(norm)
        if not code and len(norm) == 1:
            code = RIGHT_ALIASES.get(norm)
        if code and code not in codes:
            codes.append(code)

    if not codes:
        return None

    # If several rights are written, use the strongest level from the legend.
    strongest = max(codes, key=lambda code: RIGHT_ORDER.index(code) if code in RIGHT_ORDER else -1)
    return PermissionFlags(strongest, dict(RIGHT_FLAGS[strongest]))


class PermissionExcelParser:
    def __init__(self, excel_path: str):
        self.excel_path = excel_path

    def _read_matrix(self) -> Tuple[List[List[object]], int, List[Tuple[int, int]], List[Tuple[int, str]]]:
        """Читает лист одним последовательным проходом.

        Для read_only workbook нельзя многократно вызывать sheet.cell(): каждый
        случайный доступ может заново проходить XML-поток и на нескольких сотнях
        строк превращается в минуты ожидания.
        """
        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel-файл не найден: {self.excel_path}")

        workbook = load_workbook(self.excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook["2. Ролевая матрица"] if "2. Ролевая матрица" in workbook.sheetnames else workbook.active
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

        if not rows:
            raise RuntimeError("Excel-лист пуст.")
        header_index = self._find_header_index(rows)
        header = rows[header_index]

        level_columns: List[Tuple[int, int]] = []
        last_level_col = -1
        for col_idx, value in enumerate(header):
            text = clean_text(value)
            match = re.fullmatch(r"Уровень\s*(\d+)", text, flags=re.IGNORECASE)
            if match:
                level_columns.append((int(match.group(1)), col_idx))
                last_level_col = max(last_level_col, col_idx)
        level_columns.sort(key=lambda pair: pair[0])
        if not level_columns:
            raise RuntimeError("В Excel не найдены колонки вида 'Уровень 2', 'Уровень 3', ...")

        role_columns: List[Tuple[int, str]] = []
        for col_idx in range(last_level_col + 1, len(header)):
            role_name = clean_text(header[col_idx])
            if not role_name:
                continue
            if normalize_text(role_name).startswith("легенда"):
                continue
            if re.match(r"^[ПСЗР]-", role_name, flags=re.IGNORECASE):
                continue
            role_columns.append((col_idx, role_name))

        return rows, header_index, level_columns, role_columns

    def parse(self) -> List[PermissionEntry]:
        rows, header_index, level_columns, role_columns = self._read_matrix()
        if not role_columns:
            raise RuntimeError("В Excel не найдены колонки ролей после колонок уровней.")

        path_stack: List[Optional[str]] = [None] * len(level_columns)
        result: List[PermissionEntry] = []
        for row_index in range(header_index + 1, len(rows)):
            row = rows[row_index]
            for stack_idx, (_level_num, col_idx) in enumerate(level_columns):
                value = clean_text(row[col_idx] if col_idx < len(row) else None)
                if value:
                    path_stack[stack_idx] = value
                    for clear_idx in range(stack_idx + 1, len(path_stack)):
                        path_stack[clear_idx] = None

            path_parts = [part for part in path_stack if part]
            if not path_parts:
                continue

            permissions: Dict[str, PermissionFlags] = {}
            for col_idx, role_name in role_columns:
                cell_value = row[col_idx] if col_idx < len(row) else None
                flags = right_from_cell(cell_value)
                if flags:
                    permissions[role_name] = flags
            if permissions:
                result.append(
                    PermissionEntry(
                        row_number=row_index + 1,
                        path_parts=list(path_parts),
                        permissions=permissions,
                    )
                )

        if not result:
            raise RuntimeError("В Excel не найдено ни одной строки с правами.")
        return result

    def parse_paths(self) -> List[FolderPathEntry]:
        """Читает папочные строки независимо от наличия прав."""
        rows, header_index, level_columns, _role_columns = self._read_matrix()
        path_stack: List[Optional[str]] = [None] * len(level_columns)
        result: List[FolderPathEntry] = []

        for row_index in range(header_index + 1, len(rows)):
            row = rows[row_index]
            changed = False
            for stack_idx, (_level_num, col_idx) in enumerate(level_columns):
                value = clean_text(row[col_idx] if col_idx < len(row) else None)
                if not value:
                    continue
                changed = True
                path_stack[stack_idx] = value
                for clear_idx in range(stack_idx + 1, len(path_stack)):
                    path_stack[clear_idx] = None
            if not changed:
                continue
            path_parts = [part for part in path_stack if part]
            if path_parts:
                result.append(FolderPathEntry(row_number=row_index + 1, path_parts=list(path_parts)))

        if not result:
            raise RuntimeError("В Excel не найдено ни одной строки папочной структуры.")
        return result

    @staticmethod
    def _find_header_index(rows: Sequence[Sequence[object]]) -> int:
        for row_index, row in enumerate(rows[:30]):
            count = sum(
                1
                for value in row
                if re.fullmatch(r"Уровень\s*\d+", clean_text(value), flags=re.IGNORECASE)
            )
            if count >= 2:
                return row_index
        raise RuntimeError("Не найдена строка заголовков с колонками 'Уровень N'.")


# -----------------------------------------------------------------------------
# API client
# -----------------------------------------------------------------------------
class SgnlClient:
    def __init__(self, docs_url: str, hub_url: str, token: str = "", timeout: int = 40, cookies: Optional[Dict[str, str]] = None):
        self.docs_url = docs_url.strip().rstrip("/") or DEFAULT_DOCS_URL
        self.hub_url = hub_url.strip().rstrip("/") or DEFAULT_HUB_URL
        self.token = clean_text(token)
        self.timeout = timeout
        self.session = requests.Session()
        self.session.verify = VERIFY_SSL_CERTIFICATE
        self.session.trust_env = False
        if cookies:
            self.session.cookies.update({clean_text(k): clean_text(v) for k, v in cookies.items() if clean_text(k)})
        # Кэш fallback-цепочки /api/companies/find/current-user.
        # У некоторых аккаунтов /api/v1/hub/companies возвращает 403 сразу после OAuth-входа,
        # хотя список компаний доступен через маршрут текущего пользователя.
        self._company_project_ids: Dict[str, List[str]] = {}

    @property
    def headers(self) -> Dict[str, str]:
        # Для docs API нужен Bearer access_token.
        return self.docs_api_headers(include_auth=True)

    def docs_api_headers(self, include_auth: bool = True) -> Dict[str, str]:
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": self.docs_url,
            "Referer": self.docs_url + "/",
            "User-Agent": BROWSER_UA,
        }
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def hub_docs_headers(self, include_auth: bool = False, method: str = "GET") -> Dict[str, str]:
        # Точные browser-like заголовки для запросов из docs.sgnl.pro в hub.sgnl.pro.
        # В полном HAR эти hub v1 вызовы идут с Origin/Referer docs и без Content-Type на GET.
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "en-US,en;q=0.9,ru-RU;q=0.8,ru;q=0.7",
            "Origin": self.docs_url,
            "Referer": self.docs_url + "/",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-site",
            "User-Agent": BROWSER_UA,
        }
        if method.upper() != "GET":
            headers["Content-Type"] = "application/json"
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def hub_app_headers(self, include_auth: bool = True, method: str = "POST") -> Dict[str, str]:
        # Маршруты hub-приложения, которые в HAR вызываются с Origin hub.sgnl.pro.
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "en-US,en;q=0.9,ru-RU;q=0.8,ru;q=0.7",
            "Origin": self.hub_url,
            "Referer": self.hub_url + "/hub",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "same-origin",
            "User-Agent": BROWSER_UA,
        }
        if method.upper() != "GET":
            headers["Content-Type"] = "application/json"
        if include_auth:
            auth = bearer_header(self.token)
            if auth:
                headers["Authorization"] = auth
        return headers

    def _request_with_cookie_fallback(self, method: str, url: str, context: str, **kwargs) -> requests.Response:
        method_upper = method.upper()
        # 1) Как в HAR: docs -> hub, без Authorization.
        response = self.session.request(
            method,
            url,
            headers=self.hub_docs_headers(include_auth=False, method=method_upper),
            timeout=self.timeout,
            **kwargs,
        )
        if response.status_code not in (401, 403):
            return response

        # 2) Тот же контекст, но с Bearer access_token.
        bearer_response = self.session.request(
            method,
            url,
            headers=self.hub_docs_headers(include_auth=True, method=method_upper),
            timeout=self.timeout,
            **kwargs,
        )
        return bearer_response if bearer_response.status_code not in (401, 403) else response

    def _auth_origin_headers(self) -> Dict[str, str]:
        return {
            "Accept": "*/*",
            "Content-Type": "application/json",
            "Origin": "https://auth.sgnl.pro",
            "Referer": "https://auth.sgnl.pro/",
            "User-Agent": BROWSER_UA,
        }

    def _hub_form_headers(self, referer: str = "") -> Dict[str, str]:
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/x-www-form-urlencoded",
            "Origin": self.hub_url,
            "Referer": referer or (self.hub_url + "/"),
            "User-Agent": BROWSER_UA,
        }
        return headers

    def build_authorization_request(self) -> dict:
        verifier, challenge = make_pkce_pair()
        redirect_uri = self.docs_url + "/callback"
        params = {
            "client_id": OAUTH_CLIENT_ID,
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": OAUTH_SCOPE,
            "state": verifier,
            "code_challenge": challenge,
            "code_challenge_method": "S256",
        }
        return {
            "authorize_url": self.hub_url + "/connect/authorize?" + urlencode(params),
            "verifier": verifier,
            "state": verifier,
            "redirect_uri": redirect_uri,
        }

    def exchange_authorization_code(self, code: str, verifier: str, redirect_uri: str, referer: str = "") -> dict:
        code = clean_text(code)
        verifier = clean_text(verifier)
        redirect_uri = clean_text(redirect_uri)
        if not code or not verifier or not redirect_uri:
            raise RuntimeError("Недостаточно данных для OAuth token exchange.")
        response = self.session.post(
            self.hub_url + "/connect/token",
            data={
                "grant_type": "authorization_code",
                "client_id": OAUTH_CLIENT_ID,
                "redirect_uri": redirect_uri,
                "code": code,
                "code_verifier": verifier,
            },
            headers=self._hub_form_headers(referer=referer or redirect_uri),
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось получить OAuth access_token")
        data = response.json()
        token = clean_text(data.get("access_token"))
        if not token:
            raise RuntimeError(f"В ответе /connect/token нет access_token: {data}")
        self.token = token
        return data

    def login_email_password(self, email: str, password: str) -> dict:
        """Получает OAuth access_token через PKCE-flow SGNL.

        Используемые маршруты:
        - GET  /connect/authorize
        - POST /api/v1/auth/login-with-credentials
        - GET  /connect/authorize/callback
        - POST /connect/token
        """
        email = clean_text(email)
        if not email or not password:
            raise RuntimeError("Укажи email и пароль.")

        verifier, challenge = make_pkce_pair()
        redirect_uri = self.docs_url + "/callback"
        params = {
            "client_id": OAUTH_CLIENT_ID,
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": OAUTH_SCOPE,
            "state": verifier,
            "code_challenge": challenge,
            "code_challenge_method": "S256",
        }

        # 1) Создаём OAuth-запрос. Обычно будет 302 на auth.sgnl.pro.
        authorize_url = self.hub_url + "/connect/authorize?" + urlencode(params)
        response = self.session.get(authorize_url, allow_redirects=False, headers={"User-Agent": BROWSER_UA}, timeout=self.timeout)
        if response.is_redirect and response.headers.get("Location"):
            try:
                self.session.get(response.headers["Location"], allow_redirects=True, headers={"User-Agent": BROWSER_UA}, timeout=self.timeout)
            except Exception:
                # HTML-страница логина не критична для API-flow.
                pass
        elif response.status_code >= 400:
            self._raise_for_response(response, "Не удалось начать OAuth-авторизацию")

        # 2) Актуальный endpoint формы входа. На продуктиве обычно требуется
        # challengeKey от CAPTCHA, поэтому основной способ входа в GUI —
        # встроенный браузер. Этот прямой вызов оставлен только как fallback.
        login_response = self.session.post(
            self.hub_url + "/api/v1/auth/login-with-credentials",
            json={"email": email, "password": password},
            headers=self._auth_origin_headers(),
            timeout=self.timeout,
        )
        if not login_response.ok:
            body = login_response.text or ""
            if "challenge" in body.casefold() or "captcha" in body.casefold() or login_response.status_code in (400, 422):
                raise RuntimeError(
                    "Прямой вход отклонён актуальной схемой SGNL: требуется CAPTCHA/challengeKey. "
                    "Используй кнопку «Войти через браузер»."
                )
            self._raise_for_response(login_response, "Не удалось войти по email/паролю")

        # В браузерном сценарии после local-login идёт переход на hub/?postLogin=1,
        # затем открывается docs.sgnl.pro. Эти GET не обязательны для PKCE,
        # но помогают серверу воспроизвести тот же web-контекст, что в HAR.
        try:
            self.session.get(
                self.hub_url + "/?postLogin=1",
                allow_redirects=True,
                headers={"Referer": "https://auth.sgnl.pro/", "User-Agent": BROWSER_UA},
                timeout=self.timeout,
            )
        except Exception:
            pass
        try:
            self.session.get(
                self.docs_url + "/",
                allow_redirects=True,
                headers={"Referer": self.hub_url + "/", "User-Agent": BROWSER_UA},
                timeout=self.timeout,
            )
        except Exception:
            pass

        # 4) Получаем authorization code через callback.
        callback_url = self.hub_url + "/connect/authorize/callback?" + urlencode(params)
        callback_response = self.session.get(callback_url, allow_redirects=False, headers={"Referer": "https://auth.sgnl.pro/", "User-Agent": BROWSER_UA}, timeout=self.timeout)
        if callback_response.status_code not in (301, 302, 303, 307, 308):
            self._raise_for_response(callback_response, "Не удалось получить OAuth authorization code")
        location = callback_response.headers.get("Location", "")
        code = clean_text(parse_qs(urlparse(location).query).get("code", [""])[0])
        if not code:
            raise RuntimeError("OAuth callback не вернул code. Возможно, нужен CAPTCHA/2FA или изменился flow авторизации.")

        # 5) Меняем authorization code на access_token.
        return self.exchange_authorization_code(
            code=code,
            verifier=verifier,
            redirect_uri=redirect_uri,
            referer=location or callback_url,
        )

    def _raise_for_response(self, response: requests.Response, context: str) -> None:
        if response.ok:
            return
        text = response.text or ""
        if len(text) > 2500:
            text = text[:2500] + "..."
        raise RuntimeError(f"{context}. HTTP {response.status_code}: {text}")

    def get_companies(self) -> List[dict]:
        response = self._request_with_cookie_fallback("GET", f"{self.hub_url}/api/v1/hub/companies", "Не удалось загрузить компании")
        if response.status_code == 403:
            # В расширенном HAR после входа компания подтягивается так:
            # POST /api/companies/find/current-user {"deleted": false}.
            # Поэтому 403 на общем списке компаний не считаем фатальной ошибкой.
            return self.get_current_user_companies()
        self._raise_for_response(response, "Не удалось загрузить компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список компаний, получено: {type(data).__name__}")
        self._remember_company_project_ids(data)
        return data

    def get_current_user_companies(self) -> List[dict]:
        response = self.session.post(
            f"{self.hub_url}/api/companies/find/current-user",
            headers=self.hub_app_headers(include_auth=True, method="POST"),
            json={"deleted": False},
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось загрузить компании текущего пользователя")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список компаний текущего пользователя, получено: {type(data).__name__}")

        companies: List[dict] = []
        for row in data:
            if not isinstance(row, dict):
                continue
            company = row.get("company") if isinstance(row.get("company"), dict) else row
            company_id = clean_text(company.get("id"))
            if not company_id:
                continue
            item = dict(company)
            if row.get("role") and not item.get("applicationRole"):
                item["applicationRole"] = row.get("role")
            project_ids = row.get("projectIds") or row.get("project_ids") or item.get("projectIds") or []
            if isinstance(project_ids, list):
                item["projectIds"] = [clean_text(x) for x in project_ids if clean_text(x)]
                self._company_project_ids[company_id] = item["projectIds"]
            companies.append(item)
        return companies

    def _remember_company_project_ids(self, companies: Sequence[dict]) -> None:
        for company in companies:
            if not isinstance(company, dict):
                continue
            company_id = clean_text(company.get("id"))
            project_ids = company.get("projectIds") or company.get("project_ids") or []
            if company_id and isinstance(project_ids, list):
                self._company_project_ids[company_id] = [clean_text(x) for x in project_ids if clean_text(x)]

    def get_project(self, project_id: str) -> dict:
        project_id = clean_text(project_id)
        response = self._request_with_cookie_fallback(
            "POST",
            f"{self.hub_url}/api/v1/projects/get-project",
            f"Не удалось загрузить проект {project_id}",
            params={"projectId": project_id},
        )
        self._raise_for_response(response, f"Не удалось загрузить проект {project_id}")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект проекта, получено: {type(data).__name__}")
        return data

    def get_company_projects(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/projects",
            "Не удалось загрузить проекты компании",
        )
        if response.status_code == 403:
            # Fallback из /api/companies/find/current-user: там есть projectIds.
            if company_id not in self._company_project_ids:
                self.get_current_user_companies()
            project_ids = self._company_project_ids.get(clean_text(company_id), [])
            if project_ids:
                projects = []
                for project_id in project_ids:
                    try:
                        projects.append(self.get_project(project_id))
                    except Exception:
                        # Один битый/недоступный проект не должен ломать загрузку списка целиком.
                        continue
                return projects
        self._raise_for_response(response, "Не удалось загрузить проекты компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список проектов, получено: {type(data).__name__}")
        return data

    def get_custom_roles(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/custom/roles",
            "Не удалось загрузить роли компании",
        )
        self._raise_for_response(response, "Не удалось загрузить роли компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список ролей, получено: {type(data).__name__}")
        return data

    def get_company_users(self, company_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/companies/{company_id}/users",
            "Не удалось загрузить пользователей компании",
        )
        self._raise_for_response(response, "Не удалось загрузить пользователей компании")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список пользователей, получено: {type(data).__name__}")
        return data

    def get_project_roles(self, project_id: str) -> List[dict]:
        response = self._request_with_cookie_fallback(
            "GET",
            f"{self.hub_url}/api/v1/hub/projects/{project_id}/roles",
            "Не удалось загрузить роли проекта",
        )
        self._raise_for_response(response, "Не удалось загрузить роли проекта")
        data = response.json()
        if not isinstance(data, list):
            raise RuntimeError(f"Ожидался список ролей проекта, получено: {type(data).__name__}")
        return data

    def get_project_info(self, project_id: str) -> dict:
        response = self.session.get(
            f"{self.docs_url}/api/projects/{project_id}/info",
            headers=self.headers,
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось прочитать информацию о проекте")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект проекта, получено: {type(data).__name__}")
        return data

    def get_folder_tree(self, project_id: str, root_id: str = "") -> dict:
        project_id = clean_text(project_id)
        if not project_id:
            raise RuntimeError("Не указан Project ID.")
        params = {"rootId": clean_text(root_id)} if clean_text(root_id) else None
        response = self.session.get(
            f"{self.docs_url}/api/folders/project/{project_id}/tree",
            params=params,
            headers=self.docs_api_headers(include_auth=True),
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось загрузить дерево папок")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект дерева папок, получено: {type(data).__name__}")
        return data

    def create_folder(self, name: str, parent_id: str) -> str:
        name = clean_text(name)
        parent_id = clean_text(parent_id)
        if not name or not parent_id:
            raise RuntimeError("Для создания папки нужны name и parentId.")
        response = self.session.put(
            f"{self.docs_url}/api/folders",
            headers=self.docs_api_headers(include_auth=True),
            json={"name": name, "parentId": parent_id},
            timeout=self.timeout,
        )
        self._raise_for_response(response, f"Не удалось создать папку '{name}'")
        data = response.json()
        folder_id = clean_text(data.get("data") if isinstance(data, dict) else "")
        if not folder_id:
            raise RuntimeError(f"API создал папку, но не вернул её ID: {data}")
        return folder_id

    def get_permissions_tree(self, project_id: str) -> dict:
        response = self.session.get(
            f"{self.docs_url}/api/permissions/tree",
            params={"projectId": project_id},
            headers=self.headers,
            timeout=self.timeout,
        )
        self._raise_for_response(response, "Не удалось загрузить дерево прав")
        data = response.json()
        if not isinstance(data, dict):
            raise RuntimeError(f"Ожидался объект дерева, получено: {type(data).__name__}")
        return data

    def upsert_folder_permissions(self, folder_id: str, payload: List[dict]) -> None:
        response = self.session.post(
            f"{self.docs_url}/api/permissions/folders/{folder_id}/upsert",
            headers=self.headers,
            json=payload,
            timeout=self.timeout,
        )
        self._raise_for_response(response, f"Не удалось применить права для папки {folder_id}")


# -----------------------------------------------------------------------------
# Indexes / resolvers
# -----------------------------------------------------------------------------
class FolderStructureIndex:
    def __init__(self, root_id: str):
        self.root_id = clean_text(root_id)
        self.by_path: Dict[str, dict] = {"": {"id": self.root_id, "name": "Root", "path_parts": []}}

    @classmethod
    def from_tree(cls, tree: dict) -> "FolderStructureIndex":
        root_id = clean_text(tree.get("id") or tree.get("folderId"))
        if not root_id:
            raise RuntimeError("В дереве папок не найден ID корневой папки.")
        index = cls(root_id)

        def walk(node: dict, parent_parts: List[str], is_root: bool = False):
            name = clean_text(node.get("name") or node.get("folderName"))
            folder_id = clean_text(node.get("id") or node.get("folderId"))
            current_parts = list(parent_parts)
            if not is_root and name:
                current_parts.append(name)
                norm_path = normalize_path_parts(current_parts)
                if norm_path not in index.by_path:
                    index.by_path[norm_path] = {
                        "id": folder_id,
                        "name": name,
                        "path_parts": current_parts,
                    }
            for child in node.get("children") or []:
                if isinstance(child, dict):
                    walk(child, current_parts, False)

        walk(tree, [], True)
        return index

    def get(self, path_parts: Sequence[str]) -> Optional[dict]:
        return self.by_path.get(normalize_path_parts(path_parts))

    def add(self, path_parts: Sequence[str], folder_id: str):
        parts = [clean_text(part) for part in path_parts if clean_text(part)]
        self.by_path[normalize_path_parts(parts)] = {
            "id": clean_text(folder_id),
            "name": parts[-1] if parts else "Root",
            "path_parts": parts,
        }


def build_required_folder_rows(entries: Sequence[FolderPathEntry]) -> List[FolderPathEntry]:
    """Разворачивает каждую Excel-строку во все префиксы и убирает дубли."""
    result: List[FolderPathEntry] = []
    seen: set = set()
    for entry in entries:
        for depth in range(1, len(entry.path_parts) + 1):
            parts = list(entry.path_parts[:depth])
            norm_path = normalize_path_parts(parts)
            if not norm_path or norm_path in seen:
                continue
            seen.add(norm_path)
            result.append(FolderPathEntry(row_number=entry.row_number, path_parts=parts))
    return result


class FolderIndex:
    def __init__(self):
        self.items: List[FolderItem] = []
        self.path_index: Dict[str, List[FolderItem]] = {}
        self.name_index: Dict[str, List[FolderItem]] = {}

    def add(self, item: FolderItem) -> None:
        self.items.append(item)
        self.path_index.setdefault(item.norm_path, []).append(item)
        self.name_index.setdefault(normalize_text(item.name), []).append(item)

    @classmethod
    def from_tree(cls, tree: dict) -> "FolderIndex":
        index = cls()

        def node_name(node: dict) -> str:
            return clean_text(node.get("folderName") or node.get("name"))

        def node_id(node: dict) -> str:
            return clean_text(node.get("folderId") or node.get("id"))

        def walk(node: dict, parent_path: List[str], is_root: bool = False) -> None:
            name = node_name(node)
            folder_id = node_id(node)
            children = node.get("children") or []
            permissions = node.get("permissions") or []
            current_path = list(parent_path)
            if not is_root and name:
                current_path.append(name)
                if folder_id:
                    index.add(FolderItem(folder_id, name, current_path, permissions if isinstance(permissions, list) else []))
            for child in children:
                if isinstance(child, dict):
                    walk(child, current_path, is_root=False)

        walk(tree, [], is_root=True)
        return index

    def find(self, path_parts: Sequence[str], allow_suffix: bool = False) -> Tuple[Optional[FolderItem], str]:
        parts = list(path_parts)
        if parts and normalize_text(parts[0]) == "root":
            parts = parts[1:]
        norm = normalize_path_parts(parts)
        exact = self.path_index.get(norm, [])
        if len(exact) == 1:
            return exact[0], "Найдено по полному пути"
        if len(exact) > 1:
            return None, "Неоднозначный полный путь: " + "; ".join(item.path_text for item in exact[:8])

        if allow_suffix and norm:
            suffix = "/" + norm
            matches = [item for item in self.items if item.norm_path.endswith(suffix)]
            if len(matches) == 1:
                return matches[0], f"Найдено по окончанию пути: {matches[0].path_text}"
            if len(matches) > 1:
                return None, "Окончание пути неоднозначно: " + "; ".join(item.path_text for item in matches[:8])

        last = parts[-1] if parts else ""
        same_name = self.name_index.get(normalize_text(last), [])
        if same_name:
            return None, "Папка с таким именем есть, но полный путь отличается: " + "; ".join(item.path_text for item in same_name[:5])

        close = best_close_matches(norm, self.path_index.keys())
        if close:
            variants: List[str] = []
            for key in close:
                variants.extend(item.path_text for item in self.path_index.get(key, []))
            return None, "Путь не найден. Похожие: " + "; ".join(variants[:5])
        return None, "Путь не найден в дереве проекта"


class PrincipalResolver:
    def __init__(self, roles: Sequence[dict], users: Sequence[dict]):
        self.by_norm: Dict[str, Principal] = {}
        self.display_names: Dict[str, str] = {}
        self._load_roles(roles)
        self._load_users(users)

    def _register(self, key: str, principal: Principal) -> None:
        norm = normalize_text(key)
        if norm and norm not in self.by_norm:
            self.by_norm[norm] = principal
            self.display_names[norm] = principal.name

    def _load_roles(self, roles: Sequence[dict]) -> None:
        for role in roles:
            role_id = clean_text(role.get("id") or role.get("roleId"))
            name = clean_text(role.get("name"))
            if role_id and name:
                principal = Principal(name=name, id=role_id, source_type="Role")
                self._register(name, principal)

    def _load_users(self, users: Sequence[dict]) -> None:
        for user in users:
            user_id = clean_text(user.get("userId") or user.get("id"))
            email = clean_text(user.get("userEmail") or user.get("email"))
            last = clean_text(user.get("userLastName") or user.get("lastName"))
            first = clean_text(user.get("userFirstName") or user.get("firstName"))
            middle = clean_text(user.get("userMiddleName") or user.get("middleName"))
            full_name = clean_text(" ".join(part for part in [last, first, middle] if part))
            name = full_name or email
            if user_id and name:
                principal = Principal(name=name, id=user_id, source_type="User")
                self._register(name, principal)
                if email:
                    self._register(email, principal)
                if first and last:
                    self._register(f"{first} {last}", principal)

    def resolve(self, requested_name: str) -> PrincipalResolution:
        norm = normalize_text(requested_name)
        principal = self.by_norm.get(norm)
        if principal:
            return PrincipalResolution(requested_name, principal, "точное совпадение")
        close = best_close_matches(norm, self.by_norm.keys(), limit=5, cutoff=0.72)
        if close:
            variants = ", ".join(self.display_names.get(key, key) for key in close)
            return PrincipalResolution(requested_name, message=f"Не найдено точно. Похожие: {variants}")
        return PrincipalResolution(requested_name, message="Роль/пользователь не найден")


# -----------------------------------------------------------------------------
# Workers
# -----------------------------------------------------------------------------
class OAuthTokenWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, code: str, verifier: str, redirect_uri: str, callback_url: str):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.code = code
        self.verifier = verifier
        self.redirect_uri = redirect_uri
        self.callback_url = callback_url

    @Slot()
    def run(self):
        try:
            self.log.emit("OAuth code получен, выполняю обмен на access_token\n")
            client = SgnlClient(self.docs_url, self.hub_url)
            data = client.exchange_authorization_code(
                code=self.code,
                verifier=self.verifier,
                redirect_uri=self.redirect_uri,
                referer=self.callback_url,
            )
            token = clean_text(data.get("access_token"))
            payload = decode_jwt_payload(token)
            if payload:
                client_id = clean_text(payload.get("client_id"))
                scopes = payload.get("scope") or []
                scopes_text = scopes if isinstance(scopes, str) else ", ".join(clean_text(x) for x in scopes if clean_text(x))
                self.log.emit(f"  token client_id: {client_id or 'не определён'}\n")
                self.log.emit(f"  token scopes: {scopes_text or 'не определены'}\n")
                if client_id and client_id != OAUTH_CLIENT_ID:
                    self.log.emit(f"  WARNING: ожидался client_id={OAUTH_CLIENT_ID}\n")
            self.done.emit({"token": token, "raw": data, "cookies": {}})
        except Exception:
            self.failed.emit(traceback.format_exc())


class LoginWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, email: str, password: str):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.email = email
        self.password = password

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            self._log("Авторизация через SGNL OAuth\n")
            client = SgnlClient(self.docs_url, self.hub_url)
            data = client.login_email_password(self.email, self.password)
            token = clean_text(data.get("access_token"))
            payload = decode_jwt_payload(token)
            self._log("  access_token получен\n")
            if payload:
                client_id = clean_text(payload.get("client_id"))
                scopes = payload.get("scope") or []
                if isinstance(scopes, str):
                    scopes_text = scopes
                elif isinstance(scopes, list):
                    scopes_text = ", ".join(clean_text(x) for x in scopes if clean_text(x))
                else:
                    scopes_text = ""
                self._log(f"  token client_id: {client_id or 'не определён'}\n")
                self._log(f"  token scopes: {scopes_text or 'не определены'}\n")
            self.done.emit({"token": token, "raw": data, "cookies": requests.utils.dict_from_cookiejar(client.session.cookies)})
        except Exception:
            self.failed.emit(traceback.format_exc())


class LoadCatalogsWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, company_id: str = "", cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.company_id = company_id
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            self._log("Загрузка компаний и проектов\n")
            companies = client.get_companies()
            selected_company_id = clean_text(self.company_id)
            if not selected_company_id and companies:
                selected_company_id = clean_text(companies[0].get("id"))
            projects = client.get_company_projects(selected_company_id) if selected_company_id else []
            self._log(f"  компаний: {len(companies)}; проектов: {len(projects)}\n")
            self.done.emit({"companies": companies, "projects": projects, "company_id": selected_company_id})
        except Exception:
            self.failed.emit(traceback.format_exc())


class CheckAccessWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, company_id: str, cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.company_id = company_id
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            self._log("Проверка API-доступа\n")
            project = client.get_project_info(self.project_id)
            self._log(f"  проект: {clean_text(project.get('name') or project.get('title') or self.project_id)}\n")
            try:
                roles = client.get_custom_roles(self.company_id)
                self._log(f"  ролей компании: {len(roles)}\n")
            except Exception as exc:
                self._log(f"  WARNING: роли компании не загружены: {exc}\n")
                roles = client.get_project_roles(self.project_id)
                self._log(f"  ролей проекта: {len(roles)}\n")
            tree = client.get_permissions_tree(self.project_id)
            index = FolderIndex.from_tree(tree)
            self._log(f"  папок в дереве прав: {len(index.items)}\n")
            self.done.emit({"roles": len(roles), "folders": len(index.items)})
        except Exception:
            self.failed.emit(traceback.format_exc())


class PreviewWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(
        self,
        docs_url: str,
        hub_url: str,
        token: str,
        project_id: str,
        company_id: str,
        excel_path: str,
        allow_suffix: bool = False,
        limit: int = 0,
        cookies: Optional[Dict[str, str]] = None,
    ):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.company_id = company_id
        self.excel_path = excel_path
        self.allow_suffix = allow_suffix
        self.limit = limit
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            self._log("=" * 72 + "\n")
            self._log("Проверка Excel-матрицы и API SGNL\n")
            self._log(f"Excel: {self.excel_path}\n")
            self._log(f"Project ID: {self.project_id}\n")
            self._log(f"Company ID: {self.company_id}\n")
            self._log("=" * 72 + "\n\n")

            self._log("[1/4] Чтение Excel\n")
            entries = PermissionExcelParser(self.excel_path).parse()
            if self.limit > 0:
                entries = entries[: self.limit]
                self._log(f"  ограничение: первые {self.limit} строк с правами\n")
            self._log(f"  строк с правами: {len(entries)}\n")
            roles_from_excel = sorted({role for entry in entries for role in entry.permissions.keys()}, key=normalize_text)
            self._log("  роли из Excel: " + ", ".join(roles_from_excel) + "\n")

            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)

            self._log("[2/4] Загрузка ролей/пользователей\n")
            try:
                roles = client.get_custom_roles(self.company_id)
            except Exception as exc:
                roles = []
                self._log(f"  WARNING: роли компании не загружены: {exc}\n")
            try:
                project_roles = client.get_project_roles(self.project_id)
                # Merge project roles with company roles without duplicates.
                existing_ids = {clean_text(item.get("id") or item.get("roleId")) for item in roles}
                for role in project_roles:
                    rid = clean_text(role.get("id") or role.get("roleId"))
                    if rid and rid not in existing_ids:
                        roles.append(role)
                        existing_ids.add(rid)
            except Exception as exc:
                self._log(f"  WARNING: роли проекта отдельно не загружены: {exc}\n")
            try:
                users = client.get_company_users(self.company_id)
            except Exception as exc:
                users = []
                self._log(f"  WARNING: пользователи компании не загружены: {exc}\n")
            self._log(f"  ролей: {len(roles)}; пользователей: {len(users)}\n")
            resolver = PrincipalResolver(roles, users)

            self._log("[3/4] Загрузка дерева прав\n")
            tree = client.get_permissions_tree(self.project_id)
            index = FolderIndex.from_tree(tree)
            self._log(f"  папок: {len(index.items)}\n")
            self._log(f"  уникальных путей: {len(index.path_index)}\n")

            self._log("[4/4] Сопоставление путей и ролей\n")
            principal_cache: Dict[str, PrincipalResolution] = {}
            plan: List[PlanRow] = []
            for entry in entries:
                folder, folder_message = index.find(entry.path_parts, allow_suffix=self.allow_suffix)
                status = "OK" if folder else "Папка не найдена"
                principal_results: Dict[str, PrincipalResolution] = {}
                messages = [folder_message]
                missing: List[str] = []
                for role_name in entry.permissions.keys():
                    if role_name not in principal_cache:
                        principal_cache[role_name] = resolver.resolve(role_name)
                    res = principal_cache[role_name]
                    principal_results[role_name] = res
                    if not res.ok:
                        missing.append(role_name)

                if missing:
                    status = "Ошибка сопоставления"
                    messages.append("Не найдены роли/пользователи: " + ", ".join(missing[:10]))
                    for name in missing[:5]:
                        msg = principal_results[name].message
                        if msg:
                            messages.append(f"{name}: {msg}")

                plan.append(PlanRow(entry=entry, folder=folder, status=status, message=" ".join(m for m in messages if m), principals=principal_results))

            ok_count = sum(1 for row in plan if row.can_apply)
            error_count = len(plan) - ok_count
            self._log(f"Предпросмотр готов. OK: {ok_count}; с ошибками: {error_count}; всего: {len(plan)}\n")
            self.done.emit({"plan": plan, "folder_index": index})
        except Exception:
            self.failed.emit(traceback.format_exc())


class FolderStructurePreviewWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, excel_path: str, limit: int = 0, cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.excel_path = excel_path
        self.limit = limit
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            self.log.emit("=" * 72 + "\n")
            self.log.emit("Предпросмотр создания папочной структуры\n")
            self.log.emit(f"Excel: {self.excel_path}\n")
            self.log.emit(f"Project ID: {self.project_id}\n")
            self.log.emit("=" * 72 + "\n\n")

            source_rows = PermissionExcelParser(self.excel_path).parse_paths()
            if self.limit > 0:
                source_rows = source_rows[: self.limit]
                self.log.emit(f"Ограничение: первые {self.limit} папочных строк Excel\n")
            required = build_required_folder_rows(source_rows)
            self.log.emit(f"Уникальных требуемых папок: {len(required)}\n")

            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            tree = client.get_folder_tree(self.project_id)
            index = FolderStructureIndex.from_tree(tree)
            self.log.emit(f"Папок уже существует: {max(0, len(index.by_path) - 1)}\n")

            available_paths = set(index.by_path.keys())
            plan: List[FolderCreatePlanRow] = []
            for entry in required:
                norm_path = normalize_path_parts(entry.path_parts)
                parent_parts = entry.path_parts[:-1]
                parent_norm = normalize_path_parts(parent_parts)
                existing = index.by_path.get(norm_path)
                if existing:
                    actual_name = clean_text(existing.get("name"))
                    message = "Уже есть в проекте"
                    if actual_name and actual_name != entry.path_parts[-1]:
                        message += f"; фактическое имя: {actual_name}"
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Существует",
                        parent_path=display_path(parent_parts) or "Root",
                        folder_id=clean_text(existing.get("id")),
                        message=message,
                    ))
                    available_paths.add(norm_path)
                elif parent_norm in available_paths:
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Создать",
                        parent_path=display_path(parent_parts) or "Root",
                        message="Папка отсутствует и будет создана",
                    ))
                    available_paths.add(norm_path)
                else:
                    plan.append(FolderCreatePlanRow(
                        row_number=entry.row_number,
                        path_parts=list(entry.path_parts),
                        status="Ошибка",
                        parent_path=display_path(parent_parts) or "Root",
                        message="Не найден родительский путь",
                    ))

            to_create = sum(1 for row in plan if row.status == "Создать")
            existing_count = sum(1 for row in plan if row.status == "Существует")
            errors = sum(1 for row in plan if row.status == "Ошибка")
            self.log.emit(f"Итог предпросмотра: создать {to_create}; уже есть {existing_count}; ошибок {errors}\n")
            self.done.emit({"plan": plan, "to_create": to_create, "existing": existing_count, "errors": errors})
        except Exception:
            self.failed.emit(traceback.format_exc())


class CreateFoldersWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, project_id: str, plan: Sequence[FolderCreatePlanRow], cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.project_id = project_id
        self.plan = list(plan)
        self.cookies = cookies or {}

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            tree = client.get_folder_tree(self.project_id)
            index = FolderStructureIndex.from_tree(tree)
            stats = {"created": 0, "existing": 0, "failed": 0, "total": len(self.plan)}
            self.log.emit("=" * 72 + "\n")
            self.log.emit("Создание папочной структуры\n")
            self.log.emit("=" * 72 + "\n\n")

            for number, row in enumerate(self.plan, 1):
                norm_path = normalize_path_parts(row.path_parts)
                current = index.by_path.get(norm_path)
                if current:
                    stats["existing"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] SKIP: {row.path_text} — уже существует\n")
                    continue
                if row.status == "Ошибка":
                    stats["failed"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — {row.message}\n")
                    continue

                parent_parts = row.path_parts[:-1]
                parent = index.get(parent_parts)
                parent_id = clean_text(parent.get("id")) if parent else ""
                if not parent_id:
                    stats["failed"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — родитель ещё не создан\n")
                    continue

                try:
                    folder_id = client.create_folder(row.name, parent_id)
                    index.add(row.path_parts, folder_id)
                    stats["created"] += 1
                    self.log.emit(f"[{number}/{len(self.plan)}] CREATE: {row.path_text} ({folder_id})\n")
                except Exception as exc:
                    # За время между preview и apply папку мог создать другой пользователь.
                    try:
                        refreshed = FolderStructureIndex.from_tree(client.get_folder_tree(self.project_id))
                        existing = refreshed.get(row.path_parts)
                    except Exception:
                        refreshed = None
                        existing = None
                    if existing:
                        index = refreshed
                        stats["existing"] += 1
                        self.log.emit(f"[{number}/{len(self.plan)}] SKIP: {row.path_text} — появилась параллельно\n")
                    else:
                        stats["failed"] += 1
                        self.log.emit(f"[{number}/{len(self.plan)}] ERROR: {row.path_text} — {exc}\n")

            self.log.emit("\nИтог\n")
            self.log.emit(f"  создано: {stats['created']}\n")
            self.log.emit(f"  уже существовало: {stats['existing']}\n")
            self.log.emit(f"  ошибок: {stats['failed']}\n")
            self.done.emit(stats)
        except Exception:
            self.failed.emit(traceback.format_exc())


class ApplyWorker(QtCore.QObject):
    log = Signal(str)
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, docs_url: str, hub_url: str, token: str, plan: Sequence[PlanRow], cookies: Optional[Dict[str, str]] = None):
        super().__init__()
        self.docs_url = docs_url
        self.hub_url = hub_url
        self.token = token
        self.plan = list(plan)
        self.cookies = cookies or {}

    def _log(self, text: str):
        self.log.emit(text)

    @Slot()
    def run(self):
        try:
            client = SgnlClient(self.docs_url, self.hub_url, self.token, cookies=self.cookies)
            rows = [row for row in self.plan if row.can_apply]
            stats = {"folders": 0, "updated_entries": 0, "failed": 0, "skipped": len(self.plan) - len(rows)}
            self._log("=" * 72 + "\n")
            self._log("Применение прав SGNL по Excel\n")
            self._log(f"К обработке: {len(rows)}; пропущено: {stats['skipped']}\n")
            self._log("=" * 72 + "\n\n")

            for idx, row in enumerate(rows, 1):
                folder = row.folder
                assert folder is not None
                stats["folders"] += 1
                self._log(f"[{idx}/{len(rows)}] Excel строка {row.entry.row_number}: {row.entry.path_text}\n")
                self._log(f"  SGNL: {folder.path_text}\n")
                self._log(f"  folderId: {folder.id}\n")
                try:
                    payload = self._build_payload(row)
                    client.upsert_folder_permissions(folder.id, payload)
                    stats["updated_entries"] += len(payload)
                    for item in payload:
                        principal_name = self._principal_name_by_payload(row, item)
                        flags_text = self._flags_to_text(item)
                        self._log(f"  + {principal_name}: {flags_text}\n")
                except Exception as exc:
                    stats["failed"] += 1
                    self._log(f"  ERROR: {exc}\n")

            self._log("\nИтог\n")
            self._log(f"  обработано папок: {stats['folders']}\n")
            self._log(f"  отправлено записей прав: {stats['updated_entries']}\n")
            self._log(f"  ошибок: {stats['failed']}\n")
            self._log(f"  пропущено: {stats['skipped']}\n")
            self.done.emit(stats)
        except Exception:
            self.failed.emit(traceback.format_exc())

    def _build_payload(self, row: PlanRow) -> List[dict]:
        folder = row.folder
        assert folder is not None
        existing_by_source = {}
        for item in folder.permissions or []:
            source_id = clean_text(item.get("sourceId"))
            if source_id:
                existing_by_source[source_id] = item

        payload: List[dict] = []
        for role_name, perm in row.entry.permissions.items():
            res = row.principals[role_name]
            if not res.principal:
                continue
            source_id = res.principal.id
            existing = existing_by_source.get(source_id, {})
            item = {
                "read": bool(perm.flags.get("read")),
                "download": bool(perm.flags.get("download")),
                "create": bool(perm.flags.get("create")),
                "update": bool(perm.flags.get("update")),
                "delete": bool(perm.flags.get("delete")),
                "sourceId": source_id,
                "sourceType": res.principal.source_type,
            }
            existing_id = clean_text(existing.get("id"))
            if existing_id:
                item["id"] = existing_id
            payload.append(item)
        return payload

    def _principal_name_by_payload(self, row: PlanRow, payload_item: dict) -> str:
        source_id = clean_text(payload_item.get("sourceId"))
        for role_name, res in row.principals.items():
            if res.principal and res.principal.id == source_id:
                return res.principal.name
        return source_id

    @staticmethod
    def _flags_to_text(item: dict) -> str:
        parts = []
        for key, label in [("read", "read"), ("download", "download"), ("create", "create"), ("update", "update"), ("delete", "delete")]:
            if item.get(key):
                parts.append(label)
        return ", ".join(parts) if parts else "нет прав"


# -----------------------------------------------------------------------------
# UI
# -----------------------------------------------------------------------------
class ModernSection(QtWidgets.QWidget):
    def __init__(self, icon: str, title: str, collapsible: bool = True, parent=None):
        super().__init__(parent)
        self._collapsible = collapsible
        self._expanded = True
        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self.card = QtWidgets.QFrame(self)
        self.card.setObjectName("card")
        card_layout = QtWidgets.QVBoxLayout(self.card)
        card_layout.setContentsMargins(0, 0, 0, 0)
        card_layout.setSpacing(0)
        self.header = QtWidgets.QFrame(self.card)
        self.header.setObjectName("sectionHeader")
        header_layout = QtWidgets.QHBoxLayout(self.header)
        header_layout.setContentsMargins(12, 8, 12, 6)
        header_layout.setSpacing(8)
        self.icon_label = QtWidgets.QLabel(icon)
        self.icon_label.setObjectName("sectionIcon")
        self.title_label = QtWidgets.QLabel(title)
        self.title_label.setObjectName("sectionTitle")
        self.arrow_label = QtWidgets.QLabel("⌄" if collapsible else "")
        self.arrow_label.setObjectName("sectionArrow")
        header_layout.addWidget(self.icon_label)
        header_layout.addWidget(self.title_label)
        header_layout.addStretch(1)
        header_layout.addWidget(self.arrow_label)
        self.body = QtWidgets.QFrame(self.card)
        self.body.setObjectName("sectionBody")
        self.grid = QtWidgets.QGridLayout(self.body)
        self.grid.setContentsMargins(12, 0, 12, 10)
        self.grid.setHorizontalSpacing(10)
        self.grid.setVerticalSpacing(7)
        card_layout.addWidget(self.header)
        card_layout.addWidget(self.body)
        layout.addWidget(self.card)
        if collapsible:
            self.header.setCursor(QtCore.Qt.PointingHandCursor)
            self.header.mousePressEvent = self._toggle_from_mouse

    def _toggle_from_mouse(self, event):
        self.set_expanded(not self._expanded)
        event.accept()

    def set_expanded(self, expanded: bool):
        if not self._collapsible:
            return
        self._expanded = expanded
        self.body.setVisible(expanded)
        self.arrow_label.setText("⌄" if expanded else "›")


class LogDialog(QtWidgets.QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Log")
        self.resize(920, 560)
        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        self.text = QtWidgets.QTextEdit()
        self.text.setObjectName("logArea")
        self.text.setReadOnly(True)
        self.text.setFont(QtGui.QFont("Cascadia Code", 10))
        layout.addWidget(self.text, 1)

    def set_log_text(self, value: str):
        self.text.setPlainText(value)
        self.text.moveCursor(QtGui.QTextCursor.End)
        self.text.ensureCursorVisible()

    def append(self, value: str):
        self.text.moveCursor(QtGui.QTextCursor.End)
        self.text.insertPlainText(value)
        self.text.ensureCursorVisible()


class OAuthBrowserDialog(QtWidgets.QDialog):
    def __init__(self, oauth_request: dict, email: str = "", password: str = "", parent=None):
        super().__init__(parent)
        if not WEBENGINE_AVAILABLE or QtWebEngineWidgets is None:
            raise RuntimeError("Qt WebEngine недоступен.")
        self.oauth_request = dict(oauth_request)
        self.email = email
        self.password = password
        self.code = ""
        self.callback_url = ""
        self.error_text = ""

        self.setWindowTitle("Вход SGNL")
        self.resize(980, 760)
        layout = QtWidgets.QVBoxLayout(self)
        hint = QtWidgets.QLabel(
            "Выполни вход на официальной странице SGNL. Если появится CAPTCHA, реши её здесь. "
            "После успешного входа окно закроется автоматически."
        )
        hint.setWordWrap(True)
        layout.addWidget(hint)
        self.view = QtWebEngineWidgets.QWebEngineView(self)
        layout.addWidget(self.view, 1)
        buttons = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Cancel)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        self.view.urlChanged.connect(self._on_url_changed)
        self.view.loadFinished.connect(self._on_load_finished)
        self.view.setUrl(QtCore.QUrl(self.oauth_request["authorize_url"]))

    def _on_url_changed(self, qurl):
        url = qurl.toString()
        parsed = urlparse(url)
        expected = urlparse(self.oauth_request["redirect_uri"])
        if normalize_text(parsed.hostname or "") != normalize_text(expected.hostname or ""):
            return
        if not parsed.path.startswith(expected.path):
            return
        query = parse_qs(parsed.query)
        error = clean_text(query.get("error", [""])[0])
        if error:
            self.error_text = error + ": " + clean_text(query.get("error_description", [""])[0])
            self.reject()
            return
        code = clean_text(query.get("code", [""])[0])
        state = clean_text(query.get("state", [""])[0])
        if not code:
            return
        if state and state != clean_text(self.oauth_request.get("state")):
            self.error_text = "OAuth state не совпал; вход отменён из соображений безопасности."
            self.reject()
            return
        self.code = code
        self.callback_url = url
        self.accept()

    def _on_load_finished(self, ok: bool):
        if not ok or (not self.email and not self.password):
            return
        host = self.view.url().host().casefold()
        if host != "auth.sgnl.pro":
            return
        email_json = json.dumps(self.email)
        password_json = json.dumps(self.password)
        script = f"""
        (() => {{
          const email = document.querySelector('input[type="email"], input[name="email"], input[autocomplete="username"]');
          const password = document.querySelector('input[type="password"], input[name="password"], input[autocomplete="current-password"]');
          const setValue = (element, value) => {{
            if (!element || !value) return;
            const prototype = Object.getPrototypeOf(element);
            const descriptor = Object.getOwnPropertyDescriptor(prototype, 'value');
            if (descriptor && descriptor.set) descriptor.set.call(element, value);
            else element.value = value;
            element.dispatchEvent(new Event('input', {{bubbles:true}}));
            element.dispatchEvent(new Event('change', {{bubbles:true}}));
          }};
          setValue(email, {email_json});
          setValue(password, {password_json});
        }})();
        """
        try:
            self.view.page().runJavaScript(script)
        except Exception:
            pass


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("SGNL / СОД — папочная структура и права из Excel")
        self.resize(1240, 780)
        self.setMinimumSize(1080, 660)

        self.current_plan: List[PlanRow] = []
        self.current_folder_plan: List[FolderCreatePlanRow] = []
        self.active_threads: List[QtCore.QThread] = []
        self.active_workers: List[QtCore.QObject] = []
        self._log_lines: List[str] = []
        self.log_dialog: Optional[LogDialog] = None
        self._apply_after_preview = False
        self._loaded_token = ""
        self._session_cookies: Dict[str, str] = {}

        self._setup_ui()
        self._apply_stylesheet()

    def _setup_ui(self):
        self.setFont(QtGui.QFont("Segoe UI", 9))
        central = QtWidgets.QWidget()
        central.setObjectName("page")
        self.setCentralWidget(central)
        root = QtWidgets.QVBoxLayout(central)
        root.setContentsMargins(18, 14, 18, 0)
        root.setSpacing(8)

        title = QtWidgets.QLabel("SGNL / СОД — папочная структура и права из Excel")
        title.setObjectName("mainTitle")
        subtitle = QtWidgets.QLabel("Один Excel используется для создания отсутствующих папок и последующей выдачи ролевых прав.")
        subtitle.setObjectName("subtitle")
        root.addWidget(title)
        root.addWidget(subtitle)

        sec_auth = ModernSection("■", "Подключение")
        grid = sec_auth.grid
        self.ed_docs_url = QtWidgets.QLineEdit(DEFAULT_DOCS_URL)
        self.ed_hub_url = QtWidgets.QLineEdit(DEFAULT_HUB_URL)

        self.ed_email = QtWidgets.QLineEdit()
        self.ed_email.setPlaceholderText("email SGNL — необязательно, для автозаполнения браузера")
        self.ed_password = QtWidgets.QLineEdit()
        self.ed_password.setPlaceholderText("пароль — необязательно, для автозаполнения браузера")
        self.ed_password.setEchoMode(QtWidgets.QLineEdit.Password)
        self.btn_login = QtWidgets.QPushButton("Войти через браузер")
        self.btn_login.setObjectName("primaryButton")

        self.cmb_company = QtWidgets.QComboBox()
        self.cmb_company.setEditable(True)
        self.cmb_company.setPlaceholderText("Сначала войдите")
        if self.cmb_company.lineEdit():
            self.cmb_company.lineEdit().setPlaceholderText("Сначала войдите")
        self.cmb_project = QtWidgets.QComboBox()
        self.cmb_project.setEditable(True)
        self.cmb_project.setPlaceholderText("Сначала выберите компанию")
        if self.cmb_project.lineEdit():
            self.cmb_project.lineEdit().setPlaceholderText("Сначала выберите компанию")
        self.btn_load_lists = QtWidgets.QPushButton("Загрузить список")
        self.btn_load_lists.setObjectName("secondaryButton")

        # Системные ID API нужны, но в интерфейсе они скрыты.
        self.ed_project_id = QtWidgets.QLineEdit("")
        self.ed_project_id.setVisible(False)
        self.ed_company_id = QtWidgets.QLineEdit("")
        self.ed_company_id.setVisible(False)

        self.ed_token = QtWidgets.QLineEdit()
        self.ed_token.setPlaceholderText("access_token будет получен после OAuth-входа. Можно вставить вручную для диагностики.")
        self.ed_token.setEchoMode(QtWidgets.QLineEdit.Password)
        self.btn_toggle_token = QtWidgets.QToolButton()
        self.btn_toggle_token.setText("Показать")
        self.btn_toggle_token.setObjectName("eyeButton")
        self.btn_check = QtWidgets.QPushButton("Проверить доступ")
        self.btn_check.setObjectName("secondaryActionButton")
        self.lbl_auth = QtWidgets.QLabel("Доступ не проверен")
        self.lbl_auth.setObjectName("badStatus")

        grid.addWidget(self._field_label("Docs URL"), 0, 0)
        grid.addWidget(self.ed_docs_url, 0, 1, 1, 2)
        grid.addWidget(self._field_label("Hub URL"), 0, 3)
        grid.addWidget(self.ed_hub_url, 0, 4, 1, 2)

        grid.addWidget(self._field_label("Email"), 1, 0)
        grid.addWidget(self.ed_email, 1, 1, 1, 2)
        grid.addWidget(self._field_label("Пароль"), 1, 3)
        grid.addWidget(self.ed_password, 1, 4)
        grid.addWidget(self.btn_login, 1, 5)

        grid.addWidget(self._field_label("Компания"), 2, 0)
        grid.addWidget(self.cmb_company, 2, 1, 1, 2)
        grid.addWidget(self._field_label("Проект"), 2, 3)
        grid.addWidget(self.cmb_project, 2, 4)
        grid.addWidget(self.btn_load_lists, 2, 5)

        grid.addWidget(self._field_label("Token"), 3, 0)
        token_box = QtWidgets.QWidget()
        token_layout = QtWidgets.QHBoxLayout(token_box)
        token_layout.setContentsMargins(0, 0, 0, 0)
        token_layout.addWidget(self.ed_token, 1)
        token_layout.addWidget(self.btn_toggle_token)
        grid.addWidget(token_box, 3, 1, 1, 5)

        grid.addWidget(self.btn_check, 4, 1, 1, 2)
        grid.addWidget(self.lbl_auth, 4, 3, 1, 3)
        grid.setColumnStretch(1, 1)
        grid.setColumnStretch(2, 1)
        grid.setColumnStretch(4, 1)
        grid.setColumnStretch(5, 1)
        root.addWidget(sec_auth)

        sec_excel = ModernSection("■", "Excel и параметры")
        grid = sec_excel.grid
        self.ed_excel = QtWidgets.QLineEdit()
        self.ed_excel.setPlaceholderText("Выберите Excel-файл с листом '2. Ролевая матрица'")
        self.btn_excel = QtWidgets.QPushButton("Обзор...")
        self.btn_excel.setObjectName("secondaryButton")
        self.cb_suffix = QtWidgets.QCheckBox("Разрешить поиск по окончанию пути, если полный путь не найден")
        self.cb_suffix.setChecked(False)
        self.sp_limit = QtWidgets.QSpinBox()
        self.sp_limit.setRange(0, 1000000)
        self.sp_limit.setValue(0)
        self.lbl_rights = QtWidgets.QLabel("П=read; С=read+download; З=read+download+create; Р=read+download+create+update+delete; '-' не трогаем")
        self.lbl_rights.setObjectName("hintLabel")
        grid.addWidget(self._field_label("Excel-файл"), 0, 0)
        grid.addWidget(self.ed_excel, 0, 1, 1, 5)
        grid.addWidget(self.btn_excel, 0, 6)
        grid.addWidget(self._field_label("Лимит строк"), 1, 0)
        grid.addWidget(self.sp_limit, 1, 1)
        grid.addWidget(self.cb_suffix, 1, 2, 1, 5)
        grid.addWidget(self.lbl_rights, 2, 1, 1, 6)
        grid.setColumnStretch(1, 1)
        grid.setColumnStretch(2, 1)
        grid.setColumnStretch(3, 1)
        grid.setColumnStretch(4, 1)
        grid.setColumnStretch(5, 1)
        root.addWidget(sec_excel)

        self.tabs = QtWidgets.QTabWidget()
        self.tabs.setObjectName("mainTabs")

        permissions_tab = QtWidgets.QWidget()
        permissions_layout = QtWidgets.QVBoxLayout(permissions_tab)
        permissions_layout.setContentsMargins(0, 8, 0, 0)
        sec_plan = ModernSection("■", "Предпросмотр прав", collapsible=False)
        plan_grid = sec_plan.grid
        self.tbl_plan = QtWidgets.QTableWidget(0, 7)
        self.tbl_plan.setHorizontalHeaderLabels(["Статус", "Строка", "Путь Excel", "Папка SGNL", "Ролей", "Права", "Комментарий"])
        self.tbl_plan.verticalHeader().setVisible(False)
        self.tbl_plan.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_plan.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_plan.setWordWrap(False)
        header = self.tbl_plan.horizontalHeader()
        header.setSectionResizeMode(0, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(3, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(4, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(5, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(6, QtWidgets.QHeaderView.Stretch)
        plan_grid.addWidget(self.tbl_plan, 0, 0)
        permissions_layout.addWidget(sec_plan, 1)

        permissions_footer = QtWidgets.QFrame()
        permissions_footer.setObjectName("footer")
        permissions_footer_layout = QtWidgets.QHBoxLayout(permissions_footer)
        permissions_footer_layout.setContentsMargins(0, 10, 0, 10)
        self.btn_log = QtWidgets.QPushButton("Лог")
        self.btn_log.setObjectName("logBtn")
        self.btn_preview = QtWidgets.QPushButton("Проверить права")
        self.btn_preview.setObjectName("secondaryActionButton")
        self.btn_apply = QtWidgets.QPushButton("Применить права")
        self.btn_apply.setObjectName("primaryButton")
        self.btn_apply.setToolTip("Если предпросмотр ещё не построен, программа сначала выполнит проверку автоматически.")
        permissions_footer_layout.addStretch(1)
        permissions_footer_layout.addWidget(self.btn_log)
        permissions_footer_layout.addWidget(self.btn_preview)
        permissions_footer_layout.addWidget(self.btn_apply)
        permissions_layout.addWidget(permissions_footer)
        self.tabs.addTab(permissions_tab, "Ролевая матрица")

        folders_tab = QtWidgets.QWidget()
        folders_layout = QtWidgets.QVBoxLayout(folders_tab)
        folders_layout.setContentsMargins(0, 8, 0, 0)
        folder_hint = QtWidgets.QLabel(
            "Структура строится по колонкам «Уровень N». Создаются только отсутствующие папки; существующие не дублируются."
        )
        folder_hint.setObjectName("hintLabel")
        folder_hint.setWordWrap(True)
        folders_layout.addWidget(folder_hint)
        sec_folders = ModernSection("■", "Предпросмотр папочной структуры", collapsible=False)
        folder_grid = sec_folders.grid
        self.tbl_folders = QtWidgets.QTableWidget(0, 6)
        self.tbl_folders.setHorizontalHeaderLabels(["Статус", "Строка", "Путь Excel", "Родитель", "Folder ID", "Комментарий"])
        self.tbl_folders.verticalHeader().setVisible(False)
        self.tbl_folders.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.tbl_folders.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.tbl_folders.setWordWrap(False)
        folder_header = self.tbl_folders.horizontalHeader()
        folder_header.setSectionResizeMode(0, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(2, QtWidgets.QHeaderView.Stretch)
        folder_header.setSectionResizeMode(3, QtWidgets.QHeaderView.Stretch)
        folder_header.setSectionResizeMode(4, QtWidgets.QHeaderView.ResizeToContents)
        folder_header.setSectionResizeMode(5, QtWidgets.QHeaderView.Stretch)
        folder_grid.addWidget(self.tbl_folders, 0, 0)
        folders_layout.addWidget(sec_folders, 1)

        folders_footer = QtWidgets.QFrame()
        folders_footer.setObjectName("footer")
        folders_footer_layout = QtWidgets.QHBoxLayout(folders_footer)
        folders_footer_layout.setContentsMargins(0, 10, 0, 10)
        self.btn_folder_log = QtWidgets.QPushButton("Лог")
        self.btn_folder_log.setObjectName("logBtn")
        self.btn_folder_preview = QtWidgets.QPushButton("Проверить структуру")
        self.btn_folder_preview.setObjectName("secondaryActionButton")
        self.btn_folder_create = QtWidgets.QPushButton("Создать папки")
        self.btn_folder_create.setObjectName("primaryButton")
        folders_footer_layout.addStretch(1)
        folders_footer_layout.addWidget(self.btn_folder_log)
        folders_footer_layout.addWidget(self.btn_folder_preview)
        folders_footer_layout.addWidget(self.btn_folder_create)
        folders_layout.addWidget(folders_footer)
        self.tabs.addTab(folders_tab, "Папочная структура")

        root.addWidget(self.tabs, 1)
        self.statusBar().showMessage("Готово")

        self.btn_login.clicked.connect(self._login)
        self.btn_load_lists.clicked.connect(self._load_catalogs)
        self.btn_toggle_token.clicked.connect(self._toggle_token)
        self.btn_check.clicked.connect(self._check_access)
        self.btn_excel.clicked.connect(self._pick_excel)
        self.btn_log.clicked.connect(self._show_log)
        self.btn_folder_log.clicked.connect(self._show_log)
        self.btn_preview.clicked.connect(self._preview)
        self.btn_apply.clicked.connect(self._apply_permissions)
        self.btn_folder_preview.clicked.connect(self._preview_folder_structure)
        self.btn_folder_create.clicked.connect(self._create_folder_structure)
        self.cmb_company.currentIndexChanged.connect(self._on_company_changed)
        self.cmb_project.currentIndexChanged.connect(self._on_project_changed)
        self.ed_excel.textChanged.connect(self._invalidate_plan)
        self.ed_project_id.textChanged.connect(self._invalidate_plan)
        self.ed_company_id.textChanged.connect(self._invalidate_plan)
        self.ed_docs_url.textChanged.connect(self._invalidate_plan)
        self.ed_hub_url.textChanged.connect(self._invalidate_plan)

    def _field_label(self, text: str) -> QtWidgets.QLabel:
        label = QtWidgets.QLabel(text)
        label.setObjectName("fieldLabel")
        return label

    def _apply_stylesheet(self):
        self.setStyleSheet(
            """
            QMainWindow { background: #ffffff; }
            QWidget#page { background: #ffffff; color: #0f172a; }
            QLabel#mainTitle { font-size: 18px; font-weight: 700; color: #0b1220; }
            QLabel#subtitle { font-size: 12px; color: #7b879f; padding-bottom: 2px; }
            QFrame#card { background: #ffffff; border: 1px solid #dbe3ef; border-radius: 8px; }
            QFrame#sectionHeader, QFrame#sectionBody { background: transparent; border: 0; }
            QLabel#sectionIcon { color: #1277f3; font-size: 13px; font-weight: 600; min-width: 16px; }
            QLabel#sectionTitle { color: #111827; font-size: 13px; font-weight: 600; }
            QLabel#sectionArrow { color: #475569; font-size: 15px; font-weight: 600; }
            QLabel#fieldLabel { color: #111827; font-size: 13px; padding-right: 4px; }
            QLabel#hintLabel { color: #667085; font-size: 12px; }
            QLineEdit, QComboBox, QSpinBox {
                min-height: 32px; border: 1px solid #dbe3ef; border-radius: 6px;
                background: #ffffff; color: #111827; padding: 0 10px; font-size: 13px;
            }
            QLineEdit:focus, QComboBox:focus, QSpinBox:focus { border: 1px solid #93bdf8; }
            QToolButton#eyeButton { border: 0; background: transparent; color: #667085; font-size: 12px; padding: 0 6px; }
            QCheckBox { spacing: 8px; font-size: 13px; color: #1f2937; }
            QTabWidget#mainTabs::pane { border: 0; background: #ffffff; }
            QTabBar::tab { background: #f3f6fa; color: #475569; padding: 9px 18px; margin-right: 4px; border-radius: 6px; }
            QTabBar::tab:selected { background: #eaf3ff; color: #0f72ed; font-weight: 600; }
            QPushButton {
                min-height: 32px; border: 1px solid transparent; border-radius: 8px;
                padding: 0 18px; font-size: 13px; font-weight: 600;
                background: #eef2f7; color: #111827;
            }
            QPushButton:hover { background: #e5ebf4; }
            QPushButton:disabled { background: #e9edf3; color: #9aa6b8; border: 1px solid #e1e7f0; }
            QPushButton#primaryButton { background: #0f72ed; color: #ffffff; border: 1px solid #0f72ed; min-width: 170px; }
            QPushButton#primaryButton:hover { background: #0c66d8; border: 1px solid #0c66d8; }
            QPushButton#secondaryButton { background: #ffffff; border: 1px solid #dbe3ef; color: #1f2937; min-width: 110px; }
            QPushButton#secondaryButton:hover { background: #f8fafc; border: 1px solid #cbd7e6; }
            QPushButton#secondaryActionButton { background: #ffffff; border: 1px solid #dbe3ef; color: #1f2937; min-width: 180px; }
            QPushButton#logBtn { background: transparent; color: #667085; border: 1px solid #d0d5dd; min-width: 70px; padding: 0 14px; font-size: 12px; }
            QLabel#goodStatus { color: #15803d; font-weight: 650; font-size: 13px; }
            QLabel#badStatus { color: #ff0000; font-weight: 520; font-size: 13px; }
            QTableWidget {
                background: #ffffff; border: 1px solid #dbe3ef; border-radius: 8px;
                gridline-color: #e6edf5; color: #111827; selection-background-color: #eaf3ff; selection-color: #0b1220;
            }
            QTextEdit#logArea {
                border: 1px solid #d0d5dd; border-radius: 8px; background-color: #1e293b;
                color: #e2e8f0; padding: 10px; font-size: 12px;
            }
            QTableWidget::item { min-height: 24px; padding: 3px; border: 0; }
            QHeaderView::section {
                background: #ffffff; color: #111827; padding: 8px 8px; border: 0;
                border-right: 1px solid #e6edf5; border-bottom: 1px solid #e6edf5;
                font-weight: 520; font-size: 13px;
            }
            QFrame#footer { background: #ffffff; border-top: 1px solid #dbe3ef; }
            QStatusBar { background: #ffffff; color: #263244; border-top: 1px solid #dbe3ef; padding-left: 12px; font-size: 11px; min-height: 24px; }
            QStatusBar::item { border: 0; }
            QDialog { background: #ffffff; }
            """
        )

    # ------------------------------------------------------------------
    # UI helpers
    # ------------------------------------------------------------------
    def _toggle_token(self):
        if self.ed_token.echoMode() == QtWidgets.QLineEdit.Password:
            self.ed_token.setEchoMode(QtWidgets.QLineEdit.Normal)
            self.btn_toggle_token.setText("Скрыть")
        else:
            self.ed_token.setEchoMode(QtWidgets.QLineEdit.Password)
            self.btn_toggle_token.setText("Показать")

    def _log(self, text: str):
        chunk = str(text)
        if chunk and not chunk.endswith("\n"):
            chunk += "\n"
        self._log_lines.append(chunk)
        if self.log_dialog and self.log_dialog.isVisible():
            self.log_dialog.append(chunk)

    def _show_log(self):
        if self.log_dialog is None:
            self.log_dialog = LogDialog(self)
        self.log_dialog.set_log_text("".join(self._log_lines) if self._log_lines else "Лог пока пуст.\n")
        self.log_dialog.show()
        self.log_dialog.raise_()
        self.log_dialog.activateWindow()

    def _ensure_log_visible(self):
        if self.log_dialog is None:
            self.log_dialog = LogDialog(self)
        if not self.log_dialog.isVisible():
            self.log_dialog.set_log_text("".join(self._log_lines) if self._log_lines else "Лог пока пуст.\n")
            self.log_dialog.show()
        self.log_dialog.raise_()
        self.log_dialog.activateWindow()

    def _set_status_label(self, text: str, ok: bool):
        self.lbl_auth.setText(text)
        self.lbl_auth.setObjectName("goodStatus" if ok else "badStatus")
        self.lbl_auth.style().unpolish(self.lbl_auth)
        self.lbl_auth.style().polish(self.lbl_auth)

    def _show_error(self, title: str, details: str):
        self._log(f"ERROR: {title}\n{details}\n")
        QtWidgets.QMessageBox.critical(self, title, details[:4000])

    def _set_busy(self, busy: bool, message: str = ""):
        widgets = [
            self.ed_docs_url,
            self.ed_hub_url,
            self.ed_email,
            self.ed_password,
            self.btn_login,
            self.cmb_company,
            self.cmb_project,
            self.btn_load_lists,
            self.ed_project_id,
            self.ed_company_id,
            self.ed_token,
            self.btn_toggle_token,
            self.btn_check,
            self.ed_excel,
            self.btn_excel,
            self.cb_suffix,
            self.sp_limit,
            self.btn_preview,
            self.btn_apply,
            self.btn_folder_preview,
            self.btn_folder_create,
        ]
        for widget in widgets:
            widget.setEnabled(not busy)
        if message:
            self.statusBar().showMessage(message)
        if busy:
            if QtWidgets.QApplication.overrideCursor() is None:
                QtWidgets.QApplication.setOverrideCursor(QtCore.Qt.WaitCursor)
        else:
            while QtWidgets.QApplication.overrideCursor() is not None:
                QtWidgets.QApplication.restoreOverrideCursor()

    def _run_worker(self, worker: QtCore.QObject):
        thread = QtCore.QThread(self)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        for signal_name in ("done", "failed", "loaded"):
            signal = getattr(worker, signal_name, None)
            if signal is not None:
                try:
                    signal.connect(thread.quit)
                except Exception:
                    pass
        thread.finished.connect(lambda: self._cleanup_thread(thread, worker))
        self.active_threads.append(thread)
        self.active_workers.append(worker)
        thread.start()

    def _cleanup_thread(self, thread: QtCore.QThread, worker: QtCore.QObject):
        worker.deleteLater()
        thread.deleteLater()
        try:
            self.active_threads.remove(thread)
        except ValueError:
            pass
        try:
            self.active_workers.remove(worker)
        except ValueError:
            pass

    def _invalidate_plan(self):
        self.current_plan = []
        self.current_folder_plan = []
        self.tbl_plan.setRowCount(0)
        if hasattr(self, "tbl_folders"):
            self.tbl_folders.setRowCount(0)

    # ------------------------------------------------------------------
    # Data / actions
    # ------------------------------------------------------------------
    def _combo_id_by_text(self, combo: QtWidgets.QComboBox, text: str) -> str:
        norm = normalize_text(text)
        for idx in range(combo.count()):
            if normalize_text(combo.itemText(idx)) == norm:
                return clean_text(combo.itemData(idx))
        return ""

    def _current_company_id(self) -> str:
        data = self.cmb_company.currentData()
        value = clean_text(data) or self._combo_id_by_text(self.cmb_company, self.cmb_company.currentText()) or clean_text(self.ed_company_id.text())
        if value:
            self.ed_company_id.setText(value)
        return value

    def _current_project_id(self) -> str:
        data = self.cmb_project.currentData()
        value = clean_text(data) or self._combo_id_by_text(self.cmb_project, self.cmb_project.currentText()) or clean_text(self.ed_project_id.text())
        if value:
            self.ed_project_id.setText(value)
        return value

    def _set_combo_value(self, combo: QtWidgets.QComboBox, name: str, item_id: str):
        name = clean_text(name) or clean_text(item_id)
        item_id = clean_text(item_id)
        if not name and not item_id:
            return
        for idx in range(combo.count()):
            same_id = item_id and clean_text(combo.itemData(idx)) == item_id
            same_name = normalize_text(combo.itemText(idx)) == normalize_text(name)
            if same_id or same_name:
                combo.setCurrentIndex(idx)
                if item_id and not clean_text(combo.itemData(idx)):
                    combo.setItemData(idx, item_id)
                return
        combo.addItem(name, item_id)
        combo.setCurrentIndex(combo.count() - 1)

    def _on_company_changed(self, *_args):
        self._current_company_id()
        self._invalidate_plan()

    def _on_project_changed(self, *_args):
        self._current_project_id()
        self._invalidate_plan()

    def _login(self):
        docs_url = clean_text(self.ed_docs_url.text())
        hub_url = clean_text(self.ed_hub_url.text())
        if not docs_url or not hub_url:
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return
        if not WEBENGINE_AVAILABLE:
            QtWidgets.QMessageBox.critical(
                self,
                "Нет компонента браузера",
                "Для актуального входа SGNL нужен Qt WebEngine, потому что авторизация использует CAPTCHA.\n\n"
                "Для PySide6 установи полный пакет: pip install -U PySide6\n"
                "Для PyQt5: pip install PyQtWebEngine\n\n"
                "Как временный вариант можно вставить access_token вручную.",
            )
            return

        client = SgnlClient(docs_url, hub_url)
        oauth_request = client.build_authorization_request()
        dialog = OAuthBrowserDialog(
            oauth_request,
            email=clean_text(self.ed_email.text()),
            password=self.ed_password.text(),
            parent=self,
        )
        result = dialog.exec() if hasattr(dialog, "exec") else dialog.exec_()
        if not result:
            if dialog.error_text:
                self._show_error("Ошибка OAuth", dialog.error_text)
            return

        self._ensure_log_visible()
        self._set_busy(True, "Получение access_token...")
        worker = OAuthTokenWorker(
            docs_url=docs_url,
            hub_url=hub_url,
            code=dialog.code,
            verifier=oauth_request["verifier"],
            redirect_uri=oauth_request["redirect_uri"],
            callback_url=dialog.callback_url,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_login_done)
        worker.failed.connect(self._on_login_failed)
        self._run_worker(worker)

    def _on_login_done(self, payload: dict):
        token = clean_text(payload.get("token"))
        if token:
            self.ed_token.setText(token)
            self._loaded_token = token
        self._session_cookies = dict(payload.get("cookies") or {})
        self._log(f"  cookies сессии сохранены: {len(self._session_cookies)}\n")
        self._set_busy(False, "Авторизация выполнена")
        self._set_status_label("Вход выполнен, загружаю компании/проекты", True)
        self._load_catalogs()

    def _on_login_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._set_status_label("Ошибка входа", False)
        self._show_error("Ошибка входа", error)

    def _load_catalogs(self):
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни OAuth-вход через браузер.")
            return
        self._ensure_log_visible()
        self._set_busy(True, "Загрузка компаний и проектов...")
        worker = LoadCatalogsWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            company_id=self._current_company_id(),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_catalogs_loaded)
        worker.failed.connect(self._on_worker_failed_generic("Ошибка загрузки компаний/проектов"))
        self._run_worker(worker)

    def _on_catalogs_loaded(self, payload: dict):
        companies = payload.get("companies") or []
        projects = payload.get("projects") or []
        selected_company_id = clean_text(payload.get("company_id"))

        self.cmb_company.blockSignals(True)
        self.cmb_company.clear()
        for company in companies:
            self.cmb_company.addItem(clean_text(company.get("name")) or clean_text(company.get("id")), clean_text(company.get("id")))
        if selected_company_id:
            for idx in range(self.cmb_company.count()):
                if clean_text(self.cmb_company.itemData(idx)) == selected_company_id:
                    self.cmb_company.setCurrentIndex(idx)
                    break
        self.cmb_company.blockSignals(False)
        self._current_company_id()

        self.cmb_project.blockSignals(True)
        self.cmb_project.clear()
        for project in projects:
            self.cmb_project.addItem(clean_text(project.get("name")) or clean_text(project.get("id")), clean_text(project.get("id")))
        if self.cmb_project.count() == 0:
            self.cmb_project.setPlaceholderText("Сначала выберите компанию")
        if self.cmb_project.lineEdit():
            self.cmb_project.lineEdit().setPlaceholderText("Сначала выберите компанию")
        self.cmb_project.blockSignals(False)
        self._current_project_id()

        self._set_busy(False, "Списки загружены")
        self._set_status_label(f"Списки загружены: компаний {len(companies)}, проектов {len(projects)}", True)

    def _on_worker_failed_generic(self, title: str):
        def handler(error: str):
            self._set_busy(False, "Ошибка")
            self._show_error(title, error)
        return handler

    def _pick_excel(self):
        path, _ = QtWidgets.QFileDialog.getOpenFileName(self, "Выберите Excel-файл", "", "Excel (*.xlsx *.xlsm)")
        if path:
            self.ed_excel.setText(path)

    def _validate_connection_fields(self) -> bool:
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return False
        if not self._current_company_id() or not self._current_project_id():
            QtWidgets.QMessageBox.warning(
                self,
                "Не выбран проект",
                "Выбери компанию и проект из списка. Если список пустой — нажми «Войти через браузер» или «Загрузить список».",
            )
            return False
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни OAuth-вход через браузер.")
            return False
        return True

    def _validate_common(self) -> bool:
        if not self._validate_connection_fields():
            return False
        excel_path = clean_text(self.ed_excel.text())
        if not excel_path or not os.path.exists(excel_path):
            QtWidgets.QMessageBox.warning(self, "Excel не найден", "Выбери существующий Excel-файл.")
            return False
        return True

    def _validate_folder_common(self) -> bool:
        if not clean_text(self.ed_docs_url.text()) or not clean_text(self.ed_hub_url.text()):
            QtWidgets.QMessageBox.warning(self, "Не хватает данных", "Заполни Docs URL и Hub URL.")
            return False
        if not self._current_project_id():
            QtWidgets.QMessageBox.warning(self, "Не выбран проект", "Выбери проект из списка.")
            return False
        if not clean_text(self.ed_token.text()):
            QtWidgets.QMessageBox.warning(self, "Нет токена", "Сначала выполни OAuth-вход через браузер.")
            return False
        excel_path = clean_text(self.ed_excel.text())
        if not excel_path or not os.path.exists(excel_path):
            QtWidgets.QMessageBox.warning(self, "Excel не найден", "Выбери существующий Excel-файл.")
            return False
        return True

    def _check_access(self):
        if not self._validate_connection_fields():
            return
        self._ensure_log_visible()
        self._set_busy(True, "Проверка доступа...")
        worker = CheckAccessWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            company_id=self._current_company_id(),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_check_done)
        worker.failed.connect(self._on_check_failed)
        self._run_worker(worker)

    def _on_check_done(self, payload: dict):
        self._set_busy(False, "Доступ проверен")
        self._set_status_label(f"API доступен: ролей {payload['roles']}, папок {payload['folders']}", True)

    def _on_check_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._set_status_label("Ошибка доступа", False)
        self._show_error("Ошибка проверки доступа", error)

    def _make_preview_worker(self) -> PreviewWorker:
        return PreviewWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            company_id=self._current_company_id(),
            excel_path=clean_text(self.ed_excel.text()),
            allow_suffix=self.cb_suffix.isChecked(),
            limit=int(self.sp_limit.value()),
            cookies=self._session_cookies,
        )

    def _start_preview(self, apply_after_preview: bool = False):
        if not self._validate_common():
            return
        self._apply_after_preview = apply_after_preview
        self._invalidate_plan()
        self._ensure_log_visible()
        self._set_busy(True, "Проверка перед выдачей прав..." if apply_after_preview else "Построение предпросмотра...")
        worker = self._make_preview_worker()
        worker.log.connect(self._log)
        worker.done.connect(self._on_preview_done)
        worker.failed.connect(self._on_preview_failed)
        self._run_worker(worker)

    def _preview(self):
        self._start_preview(apply_after_preview=False)

    def _on_preview_done(self, payload: dict):
        self.current_plan = payload["plan"]
        self._fill_plan_table(self.current_plan)
        ok_count = sum(1 for row in self.current_plan if row.can_apply)
        error_count = len(self.current_plan) - ok_count
        can_apply_all = bool(self.current_plan) and error_count == 0
        self._set_busy(False, "Предпросмотр готов")

        if error_count:
            self._log(
                f"Предпросмотр содержит ошибки. OK-строк: {ok_count}; "
                f"строк с ошибками: {error_count}.\n"
            )
            if self._apply_after_preview:
                self._apply_after_preview = False
                if ok_count <= 0:
                    QtWidgets.QMessageBox.warning(
                        self,
                        "Права не применены",
                        "Автоматическая проверка нашла ошибки, и строк без ошибок нет. Смотри таблицу предпросмотра и лог.",
                    )
                    return
                if self._confirm_partial_apply(ok_count=ok_count, error_count=error_count):
                    self._log("Пользователь подтвердил применение только строк без ошибок.\n")
                    self._start_apply_worker()
                else:
                    self._log("Применение отменено пользователем.\n")
            return

        if can_apply_all and self._apply_after_preview:
            self._apply_after_preview = False
            self._log("Проверка успешна. Начинаю выдачу прав без второго подтверждения.\n")
            self._start_apply_worker()

    def _on_preview_failed(self, error: str):
        self.current_plan = []
        self._apply_after_preview = False
        self._fill_plan_table([])
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка предпросмотра", error)

    def _fill_plan_table(self, plan: Sequence[PlanRow]):
        self.tbl_plan.setRowCount(len(plan))
        for row_idx, row in enumerate(plan):
            folder_path = row.folder.path_text if row.folder else ""
            perms_text = "; ".join(f"{role}: {perm.text}" for role, perm in row.entry.permissions.items())
            values = [
                row.status,
                str(row.entry.row_number),
                row.entry.path_text,
                folder_path,
                str(row.permission_count),
                perms_text,
                row.message,
            ]
            for col_idx, value in enumerate(values):
                item = QtWidgets.QTableWidgetItem(value)
                item.setToolTip(value)
                if row.can_apply:
                    item.setBackground(QtGui.QColor("#e7f5e7"))
                elif row.status == "Папка не найдена":
                    item.setBackground(QtGui.QColor("#fff7d6"))
                else:
                    item.setBackground(QtGui.QColor("#ffe4e6"))
                self.tbl_plan.setItem(row_idx, col_idx, item)
        self.tbl_plan.resizeRowsToContents()

    def _confirm_apply(self, checked_rows: Optional[int] = None) -> bool:
        if checked_rows is None:
            rows_text = "Количество папок будет определено после автоматической проверки."
            action_text = "Программа сначала проверит Excel, папки и роли. Если ошибок нет, отправит upsert-запросы по папкам."
        else:
            rows_text = f"Будет изменено папок: {checked_rows}."
            action_text = "Будет использован уже построенный предпросмотр. По каждой папке будут обновлены роли из Excel."
        confirm_text = (
            f"{rows_text}\n\n"
            f"{action_text}\n\n"
            "Важно: '-' и пустые ячейки не удаляют старые права, а просто не отправляются в API.\n\n"
            "Точно применить права?"
        )
        answer = QtWidgets.QMessageBox.question(
            self,
            "Подтвердить применение прав",
            confirm_text,
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        return answer == QtWidgets.QMessageBox.Yes

    def _confirm_partial_apply(self, ok_count: int, error_count: int) -> bool:
        box = QtWidgets.QMessageBox(self)
        box.setIcon(QtWidgets.QMessageBox.Warning)
        box.setWindowTitle("Есть ошибки")
        box.setText("В предпросмотре есть ошибки.")
        box.setInformativeText(
            f"Строк без ошибок: {ok_count}.\n"
            f"Строк с ошибками: {error_count}.\n\n"
            "Можно применить права только для строк без ошибок. "
            "Ошибочные строки будут пропущены и останутся без изменений.\n\n"
            "Важно: '-' и пустые ячейки не удаляют старые права, а просто не отправляются в API."
        )
        apply_button = box.addButton("Да, применить только без ошибок", QtWidgets.QMessageBox.AcceptRole)
        cancel_button = box.addButton("Отмена", QtWidgets.QMessageBox.RejectRole)
        box.setDefaultButton(cancel_button)
        if hasattr(box, "exec"):
            box.exec()
        else:  # PyQt5 fallback
            box.exec_()
        return box.clickedButton() == apply_button

    def _start_apply_worker(self):
        self._ensure_log_visible()
        self._set_busy(True, "Выдача прав...")
        worker = ApplyWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            plan=self.current_plan,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_apply_done)
        worker.failed.connect(self._on_apply_failed)
        self._run_worker(worker)

    def _apply_permissions(self):
        if not self._validate_common():
            return
        if self.current_plan:
            ok_count = sum(1 for row in self.current_plan if row.can_apply)
            error_count = len(self.current_plan) - ok_count

            if error_count:
                if ok_count <= 0:
                    QtWidgets.QMessageBox.warning(
                        self,
                        "Нет строк без ошибок",
                        "В текущем предпросмотре нет строк, которые можно применить. Исправь Excel, пути или роли и проверь снова.",
                    )
                    return
                if not self._confirm_partial_apply(ok_count=ok_count, error_count=error_count):
                    return
                self._log(
                    f"Применение запущено частично: будет обработано {ok_count}, "
                    f"будет пропущено {error_count}.\n"
                )
                self._start_apply_worker()
                return

            if not self._confirm_apply(checked_rows=len(self.current_plan)):
                return
            self._start_apply_worker()
            return

        if not self._confirm_apply(checked_rows=None):
            return
        self._start_preview(apply_after_preview=True)

    def _preview_folder_structure(self):
        if not self._validate_folder_common():
            return
        self.current_folder_plan = []
        self.tbl_folders.setRowCount(0)
        self._ensure_log_visible()
        self._set_busy(True, "Проверка папочной структуры...")
        worker = FolderStructurePreviewWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            excel_path=clean_text(self.ed_excel.text()),
            limit=int(self.sp_limit.value()),
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_folder_preview_done)
        worker.failed.connect(self._on_folder_preview_failed)
        self._run_worker(worker)

    def _on_folder_preview_done(self, payload: dict):
        self.current_folder_plan = list(payload.get("plan") or [])
        self._fill_folder_table(self.current_folder_plan)
        self._set_busy(False, "Предпросмотр структуры готов")
        if payload.get("errors"):
            QtWidgets.QMessageBox.warning(
                self,
                "Есть ошибки в структуре",
                f"К созданию: {payload.get('to_create', 0)}\n"
                f"Уже существует: {payload.get('existing', 0)}\n"
                f"Ошибок: {payload.get('errors', 0)}\n\n"
                "Ошибочные строки созданы не будут.",
            )

    def _on_folder_preview_failed(self, error: str):
        self.current_folder_plan = []
        self.tbl_folders.setRowCount(0)
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка проверки папочной структуры", error)

    def _fill_folder_table(self, plan: Sequence[FolderCreatePlanRow]):
        self.tbl_folders.setRowCount(len(plan))
        for row_idx, row in enumerate(plan):
            values = [
                row.status,
                str(row.row_number),
                row.path_text,
                row.parent_path,
                row.folder_id,
                row.message,
            ]
            for col_idx, value in enumerate(values):
                item = QtWidgets.QTableWidgetItem(value)
                item.setToolTip(value)
                if row.status == "Существует":
                    item.setBackground(QtGui.QColor("#e7f5e7"))
                elif row.status == "Создать":
                    item.setBackground(QtGui.QColor("#fff7d6"))
                else:
                    item.setBackground(QtGui.QColor("#ffe4e6"))
                self.tbl_folders.setItem(row_idx, col_idx, item)
        self.tbl_folders.resizeRowsToContents()

    def _create_folder_structure(self):
        if not self._validate_folder_common():
            return
        if not self.current_folder_plan:
            QtWidgets.QMessageBox.information(
                self,
                "Сначала проверка",
                "Сначала нажми «Проверить структуру». Это исключает создание папок не в том проекте или не под тем родителем.",
            )
            return
        create_count = sum(1 for row in self.current_folder_plan if row.status == "Создать")
        error_count = sum(1 for row in self.current_folder_plan if row.status == "Ошибка")
        if create_count == 0:
            QtWidgets.QMessageBox.information(self, "Создавать нечего", "Все папки из Excel уже существуют.")
            return
        answer = QtWidgets.QMessageBox.question(
            self,
            "Подтвердить создание папок",
            f"Будет создано отсутствующих папок: {create_count}.\n"
            f"Ошибочных строк будет пропущено: {error_count}.\n\n"
            "Существующие папки не изменяются и не дублируются. Продолжить?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer != QtWidgets.QMessageBox.Yes:
            return

        self._ensure_log_visible()
        self._set_busy(True, "Создание папок...")
        worker = CreateFoldersWorker(
            docs_url=clean_text(self.ed_docs_url.text()),
            hub_url=clean_text(self.ed_hub_url.text()),
            token=clean_text(self.ed_token.text()),
            project_id=self._current_project_id(),
            plan=self.current_folder_plan,
            cookies=self._session_cookies,
        )
        worker.log.connect(self._log)
        worker.done.connect(self._on_folder_create_done)
        worker.failed.connect(self._on_folder_create_failed)
        self._run_worker(worker)

    def _on_folder_create_done(self, stats: dict):
        self._set_busy(False, "Папочная структура создана")
        QtWidgets.QMessageBox.information(
            self,
            "Готово",
            f"Создано папок: {stats.get('created', 0)}\n"
            f"Уже существовало: {stats.get('existing', 0)}\n"
            f"Ошибок: {stats.get('failed', 0)}",
        )
        # Дерево изменилось — старый предпросмотр прав больше нельзя считать актуальным.
        self.current_plan = []
        self.tbl_plan.setRowCount(0)
        self._preview_folder_structure()

    def _on_folder_create_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка создания папочной структуры", error)

    def _on_apply_done(self, stats: dict):
        self._set_busy(False, "Готово")
        QtWidgets.QMessageBox.information(
            self,
            "Готово",
            f"Обработано папок: {stats['folders']}\n"
            f"Отправлено записей прав: {stats['updated_entries']}\n"
            f"Ошибок: {stats['failed']}\n"
            f"Пропущено: {stats['skipped']}",
        )

    def _on_apply_failed(self, error: str):
        self._set_busy(False, "Ошибка")
        self._show_error("Ошибка выдачи прав", error)


def main() -> int:
    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow()
    window.show()
    if hasattr(app, "exec"):
        return app.exec()
    return app.exec_()


if __name__ == "__main__":
    sys.exit(main())
